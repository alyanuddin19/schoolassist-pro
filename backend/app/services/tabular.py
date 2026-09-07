"""Small CSV/XLSX reader used by upload endpoints without pandas."""

import csv
import io
from typing import Any

from openpyxl import load_workbook


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    return value


def read_table(content: bytes, filename: str | None) -> tuple[list[str], list[dict[str, Any]]]:
    if (filename or "").lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        headers = [_normalize_header(column) for column in (reader.fieldnames or [])]
        rows = []
        for raw in reader:
            normalized = {
                _normalize_header(key): _clean_cell(value)
                for key, value in raw.items()
                if key is not None
            }
            rows.append(normalized)
        return headers, rows

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [_normalize_header(column) for column in raw_headers]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(
            {
                header: _clean_cell(values[index] if index < len(values) else "")
                for index, header in enumerate(headers)
                if header
            }
        )
    workbook.close()
    return headers, rows
